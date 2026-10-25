"""
Excel Export generator for CO-PO Attainment System.
Generates comprehensive multi-sheet workbooks and specialized exports using openpyxl.
Includes professional styling, clean borders, preserved original structures,
teacher metadata, student attainment, and OBE metrics.
"""

import io
from typing import Dict, Any, List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Styling Palettes (Clean Academic Minimalist)
HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
ACCENT_FILL = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
LIGHT_ACCENT_FILL = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

TITLE_FONT = Font(name="Calibri", size=15, bold=True, color="1E293B")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1E293B")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="0F172A")
REGULAR_FONT = Font(name="Calibri", size=10, color="0F172A")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
CELL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
HEADER_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=Side(border_style="medium", color="0F172A"))


def auto_fit_columns(ws, min_width=12, max_width=45):
    """Adjusts column widths based on cell content."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if "\n" in val_str:
                lines = val_str.split("\n")
                line_len = max(len(l) for l in lines)
                max_len = max(max_len, line_len)
            else:
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


def style_header_row(ws, row_idx: int, col_count: int, fill=HEADER_FILL, font=HEADER_FONT):
    """Applies clean header formatting."""
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = HEADER_BORDER
    ws.row_dimensions[row_idx].height = 28


def create_complete_excel_report(report_data: Dict[str, Any]) -> io.BytesIO:
    """
    Builds the complete multi-sheet Excel report with:
    1. Faculty & Course Info
    2. Assessment Summaries
    3. Student Marks & Attainment
    4. Direct CO Attainment
    5. Indirect Ratings & Attainment
    6. Weighted CO Attainment
    7. CO-PO Mapping Matrix
    8. PO Attainment
    9. Graph & Distribution Data
    """
    wb = Workbook()
    
    teacher = report_data.get('teacher') or {}
    direct = (report_data.get('direct') or {}).get('data') or {}
    indirect = (report_data.get('indirect') or {}).get('data') or {}
    weighted = (report_data.get('weighted') or {}).get('data') or {}
    copo = report_data.get('copo') or {}
    po_results = copo.get('po_results') or []
    mapping_matrix = copo.get('matrix') or {}

    # ========================================================
    # Sheet 1: Course & Faculty Information
    # ========================================================
    ws_info = wb.active
    ws_info.title = "Course & Faculty Info"
    ws_info.views.sheetView[0].showGridLines = True

    ws_info["A1"] = "D. Y. PATIL PRATISHTHAN'S D. Y. PATIL COLLEGE OF ENGINEERING, AKURDI, PUNE - 411044"
    ws_info["A1"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    ws_info["A2"] = "Approved by AICTE, Recognized by DTE, Affiliated to SPPU, Accredited by NAAC with 'A' Grade, NBA Accredited, Autonomous AY 2024-25"
    ws_info["A2"].font = Font(name="Calibri", size=9, italic=True, color="475569")

    if teacher.get("course_type") == "practical":
        ws_info["A4"] = "LABORATORY CO-PO ATTAINMENT EVALUATION REPORT"
    else:
        ws_info["A4"] = "CO-PO ATTAINMENT EVALUATION REPORT"
    ws_info["A4"].font = TITLE_FONT
    ws_info.merge_cells("A4:D4")
    ws_info["A4"].alignment = Alignment(horizontal="left", vertical="center")

    ws_info["A6"] = "Teacher & Academic Course Metadata"
    ws_info["A6"].font = SECTION_FONT

    # Calculate fallback overall attainment if weighted is not explicitly run
    d_summaries = direct.get('co_summaries', [])
    fallback_att = "N/A"
    fallback_lvl = "N/A"
    if d_summaries:
        valid_atts = [float(c.get('attainment', 0) or 0) for c in d_summaries if c.get('attainment') is not None]
        if valid_atts:
            avg_a = sum(valid_atts) / len(valid_atts)
            fallback_att = f"{avg_a:.2f}"
            fallback_lvl = "Level 3" if avg_a >= 2.5 else ("Level 2" if avg_a >= 1.5 else "Level 1")

    cur_overall_att = weighted.get('overall_attainment')
    if cur_overall_att is None or cur_overall_att == '':
        cur_overall_att = fallback_att
    cur_overall_lvl = weighted.get('overall_level_label')
    if not cur_overall_lvl or cur_overall_lvl == 'N/A':
        cur_overall_lvl = fallback_lvl

    course_type_str = f"Practical (Batch: {teacher.get('batch')})" if teacher.get("course_type") == "practical" and teacher.get("batch") else ("Practical" if teacher.get("course_type") == "practical" else "Theory")

    meta_rows = [
        ("Faculty Name", teacher.get("faculty_name", "N/A")),
        ("Subject Name & Code", teacher.get("subject", "N/A")),
        ("Course Nature & Batch", course_type_str),
        ("Class", teacher.get("class_name", "N/A")),
        ("Division", teacher.get("division", "N/A")),
        ("Semester", teacher.get("semester", "N/A")),
        ("Academic Year", teacher.get("academic_year", "N/A")),
        ("Direct Assessment Weight", f"{weighted.get('direct_weight', 80)}%"),
        ("Indirect Assessment Weight", f"{weighted.get('indirect_weight', 20)}%"),
        ("Overall Course Attainment", f"{cur_overall_att} / 3.00"),
        ("Overall Attainment Level", cur_overall_lvl),
    ]

    r_idx = 7
    for label, val in meta_rows:
        ws_info.cell(row=r_idx, column=1, value=label).font = BOLD_FONT
        ws_info.cell(row=r_idx, column=1).fill = LIGHT_ACCENT_FILL
        ws_info.cell(row=r_idx, column=1).border = CELL_BORDER
        
        ws_info.cell(row=r_idx, column=2, value=val).font = REGULAR_FONT
        ws_info.cell(row=r_idx, column=2).border = CELL_BORDER
        r_idx += 1

    auto_fit_columns(ws_info)

    # ========================================================
    # Sheet 2: Assessment Setup & Passing Summaries
    # ========================================================
    ws_assess = wb.create_sheet(title="Assessment Summaries")
    ws_assess.views.sheetView[0].showGridLines = True

    ws_assess["A1"] = "Direct Assessment Tools & Passing Summaries (Test/CCE/Unit Tests)"
    ws_assess["A1"].font = SECTION_FONT
    
    headers_assess = [
        "Assessment Name", "Mapped CO(s)", "Max Marks", "Passing Mark (50%)",
        "Students Appeared", "Students Passed", "Students Failed", 
        "Pass Percentage (%)", "Average Marks", "Min Marks", "Max Scored"
    ]
    for c_idx, h in enumerate(headers_assess, start=1):
        ws_assess.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_assess, 3, len(headers_assess))

    r_idx = 4
    for a in direct.get("assessment_summaries", []):
        row_vals = [
            a.get("name"), a.get("mapped_cos"), a.get("max_marks"), a.get("passing_mark"),
            a.get("appeared"), a.get("passed"), a.get("failed"),
            a.get("pass_pct"), a.get("avg_marks"), a.get("min_marks"), a.get("max_marks_scored")
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_assess.cell(row=r_idx, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx > 2 else "left")
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL
        r_idx += 1

    auto_fit_columns(ws_assess)

    # ========================================================
    # Sheet 3: Student Marks & Attainment (Preserving Table Structure)
    # ========================================================
    ws_students = wb.create_sheet(title="Student Marks & Attainment")
    ws_students.views.sheetView[0].showGridLines = True

    active_cols = direct.get("active_columns", [])
    col_names = [c["name"] for c in active_cols]

    headers_student = ["Roll No", "Student Name"] + col_names + ["Total Obtained", "Max Total", "Percentage (%)", "Passed COs"]
    for c_idx, h in enumerate(headers_student, start=1):
        ws_students.cell(row=2, column=c_idx, value=h)
    style_header_row(ws_students, 2, len(headers_student))

    r_idx = 3
    for s in direct.get("students", []):
        row_vals = [s.get("roll_no"), s.get("name")]
        for col_name in col_names:
            row_vals.append(s.get("marks", {}).get(col_name, ""))
        row_vals.extend([
            s.get("total_obtained"),
            s.get("total_max"),
            s.get("percentage"),
            ", ".join(s.get("passed_cos", []))
        ])

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_students.cell(row=r_idx, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx not in [2, len(row_vals)] else "left")
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL
        r_idx += 1

    auto_fit_columns(ws_students)

    # ========================================================
    # Sheet 4: Direct CO Attainment
    # ========================================================
    ws_d_co = wb.create_sheet(title="Direct CO Attainment")
    ws_d_co.views.sheetView[0].showGridLines = True

    ws_d_co["A1"] = "Direct Course Outcome (CO) Attainment Calculations"
    ws_d_co["A1"].font = SECTION_FONT

    headers_d_co = [
        "Course Outcome", "Mapped Assessments", "Max Alloc. Marks",
        "Students Evaluated", "Students Passed (>=50%)", "Pass Percentage (%)",
        "Internal (CCE/UT) Attainment", "ESE (External) Attainment",
        "Direct Attainment (Points)", "Attainment Level"
    ]
    for c_idx, h in enumerate(headers_d_co, start=1):
        ws_d_co.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_d_co, 3, len(headers_d_co))

    r_idx = 4
    for c in direct.get("co_summaries", []):
        int_att = c.get("internal_attainment", "-") if c.get("has_internal") else "N/A"
        ese_att = c.get("ese_attainment", "-") if c.get("has_ese") else "N/A"

        row_vals = [
            c.get("co"), c.get("mapped_assessments"), c.get("max_marks"),
            c.get("students_appeared"), c.get("students_passed"), c.get("pass_percentage"),
            int_att, ese_att,
            c.get("attainment"), c.get("level_label")
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_d_co.cell(row=r_idx, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx not in [2, 10, 12, 13] else "left")
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL
        r_idx += 1

    auto_fit_columns(ws_d_co)

    # ========================================================
    # Sheet 5: Indirect Survey & Attainment
    # ========================================================
    ws_ind = wb.create_sheet(title="Indirect CO Attainment")
    ws_ind.views.sheetView[0].showGridLines = True

    ws_ind["A1"] = "Indirect Feedback & Survey Attainment Calculations"
    ws_ind["A1"].font = SECTION_FONT

    headers_ind = [
        "Course Outcome", "Survey Responses", "Max Rating Scale",
        "Average Rating", "Rating Percentage (%)", "Attainment (0-3)", "Attainment Level"
    ]
    for c_idx, h in enumerate(headers_ind, start=1):
        ws_ind.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_ind, 3, len(headers_ind))

    r_idx = 4
    for c in indirect.get("co_summaries", []):
        row_vals = [
            c.get("co"), c.get("response_count"), c.get("max_rating"),
            c.get("average_rating"), c.get("rating_percentage"),
            c.get("attainment"), c.get("level_label")
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_ind.cell(row=r_idx, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx != 7 else "left")
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL
        r_idx += 1

    # Also list cleaned respondent ratings if available
    cleaned_rows = indirect.get("cleaned_rows", [])
    if cleaned_rows:
        r_idx += 2
        ws_ind.cell(row=r_idx, column=1, value="Respondent Indirect Feedback Ratings").font = SECTION_FONT
        r_idx += 1
        ind_cos = indirect.get("detected_cos", [])
        ind_hdr = ["Respondent #"] + ind_cos
        for c_idx, h in enumerate(ind_hdr, start=1):
            ws_ind.cell(row=r_idx, column=c_idx, value=h)
        style_header_row(ws_ind, r_idx, len(ind_hdr), fill=SUBHEADER_FILL)
        
        r_idx += 1
        for idx, row_item in enumerate(cleaned_rows, start=1):
            ws_ind.cell(row=r_idx, column=1, value=f"Resp {idx}").border = CELL_BORDER
            for c_idx, co in enumerate(ind_cos, start=2):
                cell = ws_ind.cell(row=r_idx, column=c_idx, value=row_item.get(co, ""))
                cell.font = REGULAR_FONT
                cell.border = CELL_BORDER
                cell.alignment = Alignment(horizontal="center")
            r_idx += 1

    auto_fit_columns(ws_ind)

    # ========================================================
    # Sheet 6: Weighted Combined CO Attainment
    # ========================================================
    ws_wt = wb.create_sheet(title="Weighted CO Attainment")
    ws_wt.views.sheetView[0].showGridLines = True

    ws_wt["A1"] = f"Overall Weighted CO Attainment ({weighted.get('direct_weight', 80)}% Direct + {weighted.get('indirect_weight', 20)}% Indirect)"
    ws_wt["A1"].font = SECTION_FONT

    headers_wt = [
        "Course Outcome", "Direct %", "Direct Contrib. (%)", "Direct Attainment",
        "Indirect %", "Indirect Contrib. (%)", "Indirect Attainment",
        "Combined %", "Combined Attainment (0-3)", "Combined Level"
    ]
    for c_idx, h in enumerate(headers_wt, start=1):
        ws_wt.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_wt, 3, len(headers_wt))

    r_idx = 4
    for c in weighted.get("combined_results", []):
        row_vals = [
            c.get("co"), c.get("direct_pct"), c.get("direct_contribution"), c.get("direct_attainment"),
            c.get("indirect_pct"), c.get("indirect_contribution"), c.get("indirect_attainment"),
            c.get("combined_pct"), c.get("combined_attainment"), c.get("level_label")
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_wt.cell(row=r_idx, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx != 10 else "left")
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL
        r_idx += 1

    auto_fit_columns(ws_wt)

    # ========================================================
    # Sheet 7: CO-PO Mapping Matrix
    # ========================================================
    ws_map = wb.create_sheet(title="CO-PO Mapping Matrix")
    ws_map.views.sheetView[0].showGridLines = True

    ws_map["A1"] = "CO-PO / CO-PSO Correlation Articulation Matrix"
    ws_map["A1"].font = SECTION_FONT

    po_list = copo.get("po_list", [f"PO{i}" for i in range(1, 13)])
    co_list = copo.get("co_list", [c["co"] for c in weighted.get("combined_results", [])])

    headers_map = ["Course Outcome"] + po_list
    for c_idx, h in enumerate(headers_map, start=1):
        ws_map.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_map, 3, len(headers_map))

    r_idx = 4
    for co in co_list:
        ws_map.cell(row=r_idx, column=1, value=co).font = BOLD_FONT
        ws_map.cell(row=r_idx, column=1).border = CELL_BORDER
        ws_map.cell(row=r_idx, column=1).fill = LIGHT_ACCENT_FILL
        ws_map.cell(row=r_idx, column=1).alignment = Alignment(horizontal="center")

        for c_idx, po in enumerate(po_list, start=2):
            val = mapping_matrix.get(co, {}).get(po, 0)
            cell = ws_map.cell(row=r_idx, column=c_idx, value=val if val != 0 else "-")
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center")
            if val and val > 0:
                cell.font = BOLD_FONT
        r_idx += 1

    # Reference DY Patil format: 4 summary rows under matrix
    po_map = {p.get("po"): p for p in po_results}
    summary_rows_def = [
        ("CO-PO-PSO Mapping (Average Level)", "mapping_avg", PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")),
        ("CO-PO-PSO Direct Attainment", "direct_attainment", PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")),
        ("CO-PO-PSO Indirect Attainment", "indirect_attainment", PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")),
        ("Final / Combined Attainment", "attainment", PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"))
    ]

    for label, field_key, fill_color in summary_rows_def:
        ws_map.cell(row=r_idx, column=1, value=label).font = BOLD_FONT
        ws_map.cell(row=r_idx, column=1).border = CELL_BORDER
        ws_map.cell(row=r_idx, column=1).fill = fill_color
        ws_map.cell(row=r_idx, column=1).alignment = Alignment(horizontal="left")

        for c_idx, po in enumerate(po_list, start=2):
            p_obj = po_map.get(po, {})
            val = p_obj.get(field_key)
            if val is None or val == "":
                disp_val = "-"
            else:
                disp_val = val
            cell = ws_map.cell(row=r_idx, column=c_idx, value=disp_val)
            cell.font = BOLD_FONT
            cell.border = CELL_BORDER
            cell.fill = fill_color
            cell.alignment = Alignment(horizontal="center")
        r_idx += 1

    auto_fit_columns(ws_map)

    # ========================================================
    # Sheet 8: PO Attainment Results
    # ========================================================
    ws_po = wb.create_sheet(title="PO Attainment")
    ws_po.views.sheetView[0].showGridLines = True

    ws_po["A1"] = "Program Outcome (PO & PSO) Attainment Results"
    ws_po["A1"].font = SECTION_FONT

    headers_po = [
        "Program Outcome", "Mapping Level Sum", "Average Mapping Level",
        "Direct PO Attainment", "Indirect PO Attainment",
        "PO Attainment (0-3)", "Attainment Level"
    ]
    for c_idx, h in enumerate(headers_po, start=1):
        ws_po.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_po, 3, len(headers_po))

    r_idx = 4
    for p in po_results:
        att = p.get("attainment")
        att_display = att if att is not None else "N/A"
        dir_att = p.get("direct_attainment")
        dir_display = dir_att if dir_att is not None else "N/A"
        ind_att = p.get("indirect_attainment")
        ind_display = ind_att if ind_att is not None else "N/A"
        avg_map = p.get("mapping_avg", "N/A")

        row_vals = [
            p.get("po"), p.get("mapping_sum"), avg_map,
            dir_display, ind_display,
            att_display, p.get("level_label")
        ]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws_po.cell(row=r_idx, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx != 7 else "left")
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL
        r_idx += 1

    auto_fit_columns(ws_po)

    # ========================================================
    # Sheet 9: Graph & Performance Analytics Data
    # ========================================================
    ws_graph = wb.create_sheet(title="Graph & Distribution Data")
    ws_graph.views.sheetView[0].showGridLines = True

    ws_graph["A1"] = "Chart Data - Student Performance Distribution"
    ws_graph["A1"].font = SECTION_FONT

    ws_graph["A3"] = "Score Range"
    ws_graph["B3"] = "Number of Students"
    style_header_row(ws_graph, 3, 2, fill=SUBHEADER_FILL)

    ranges = direct.get("score_ranges", {})
    r_idx = 4
    for rng_key, cnt in ranges.items():
        ws_graph.cell(row=r_idx, column=1, value=rng_key).border = CELL_BORDER
        ws_graph.cell(row=r_idx, column=2, value=cnt).border = CELL_BORDER
        r_idx += 1

    r_idx += 2
    ws_graph.cell(row=r_idx, column=1, value="Chart Data - CO Attainment Comparison").font = SECTION_FONT
    r_idx += 1

    headers_chart_co = ["Course Outcome", "Direct Attainment", "Indirect Attainment", "Combined Attainment"]
    for c_idx, h in enumerate(headers_chart_co, start=1):
        ws_graph.cell(row=r_idx, column=c_idx, value=h)
    style_header_row(ws_graph, r_idx, len(headers_chart_co), fill=SUBHEADER_FILL)
    r_idx += 1

    for c in weighted.get("combined_results", []):
        ws_graph.cell(row=r_idx, column=1, value=c.get("co")).border = CELL_BORDER
        ws_graph.cell(row=r_idx, column=2, value=c.get("direct_attainment")).border = CELL_BORDER
        ws_graph.cell(row=r_idx, column=3, value=c.get("indirect_attainment")).border = CELL_BORDER
        ws_graph.cell(row=r_idx, column=4, value=c.get("combined_attainment")).border = CELL_BORDER
        r_idx += 1

    auto_fit_columns(ws_graph)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def create_direct_only_excel(direct_data: Dict[str, Any], teacher: Dict[str, Any]) -> io.BytesIO:
    """Generates direct assessment only Excel report."""
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "Direct CO Attainment"
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum["A1"] = "D. Y. PATIL PRATISHTHAN'S D. Y. PATIL COLLEGE OF ENGINEERING, AKURDI, PUNE - 411044"
    ws_sum["A1"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    ws_sum["A2"] = "Approved by AICTE, Recognized by DTE, Affiliated to SPPU, Accredited by NAAC with 'A' Grade, NBA Accredited, Autonomous AY 2024-25"
    ws_sum["A2"].font = Font(name="Calibri", size=9, italic=True, color="475569")

    ws_sum["A4"] = f"Direct Assessment Report - {teacher.get('subject', 'Subject')}"
    ws_sum["A4"].font = TITLE_FONT

    headers = [
        "Course Outcome", "Mapped Assessments", "Max Marks", "Appeared", "Passed", "Pass %",
        "Internal (CCE/UT) Attainment", "ESE (External) Attainment",
        "Direct Attainment (Points)", "Attainment Level"
    ]
    for c_idx, h in enumerate(headers, start=1):
        ws_sum.cell(row=6, column=c_idx, value=h)
    style_header_row(ws_sum, 6, len(headers))

    r_idx = 7
    for c in direct_data.get("co_summaries", []):
        int_att = c.get("internal_attainment", "-") if c.get("has_internal") else "N/A"
        ese_att = c.get("ese_attainment", "-") if c.get("has_ese") else "N/A"

        row_vals = [
            c.get("co"), c.get("mapped_assessments"), c.get("max_marks"),
            c.get("students_appeared"), c.get("students_passed"), c.get("pass_percentage"),
            int_att, ese_att,
            c.get("attainment"), c.get("level_label")
        ]
        for c_idx, v in enumerate(row_vals, start=1):
            cell = ws_sum.cell(row=r_idx, column=c_idx, value=v)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx not in [2, 10, 12, 13] else "left")
        r_idx += 1

    auto_fit_columns(ws_sum)

    # Student marks sheet
    ws_marks = wb.create_sheet(title="Student Marks")
    ws_marks.views.sheetView[0].showGridLines = True
    active_cols = direct_data.get("active_columns", [])
    col_names = [c["name"] for c in active_cols]
    m_hdrs = ["Roll No", "Student Name"] + col_names + ["Total", "Max", "Percentage", "Passed COs"]
    for c_idx, h in enumerate(m_hdrs, start=1):
        ws_marks.cell(row=1, column=c_idx, value=h)
    style_header_row(ws_marks, 1, len(m_hdrs))

    r_idx = 2
    for s in direct_data.get("students", []):
        row_vals = [s.get("roll_no"), s.get("name")]
        for col_name in col_names:
            row_vals.append(s.get("marks", {}).get(col_name, ""))
        row_vals.extend([s.get("total_obtained"), s.get("total_max"), s.get("percentage"), ", ".join(s.get("passed_cos", []))])
        for c_idx, v in enumerate(row_vals, start=1):
            cell = ws_marks.cell(row=r_idx, column=c_idx, value=v)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx not in [2, len(row_vals)] else "left")
        r_idx += 1

    auto_fit_columns(ws_marks)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def create_indirect_only_excel(indirect_data: Dict[str, Any], teacher: Dict[str, Any]) -> io.BytesIO:
    """Generates indirect assessment only Excel report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Indirect CO Attainment"
    ws.views.sheetView[0].showGridLines = True

    ws["A1"] = "D. Y. PATIL PRATISHTHAN'S D. Y. PATIL COLLEGE OF ENGINEERING, AKURDI, PUNE - 411044"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    ws["A2"] = "Approved by AICTE, Recognized by DTE, Affiliated to SPPU, Accredited by NAAC with 'A' Grade, NBA Accredited, Autonomous AY 2024-25"
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="475569")

    ws["A4"] = f"Indirect Survey Report - {teacher.get('subject', 'Subject')}"
    ws["A4"].font = TITLE_FONT

    headers = ["Course Outcome", "Responses", "Max Scale", "Average Rating", "Rating %", "Attainment", "Level"]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=6, column=c_idx, value=h)
    style_header_row(ws, 6, len(headers))

    r_idx = 7
    for c in indirect_data.get("co_summaries", []):
        row_vals = [
            c.get("co"), c.get("response_count"), c.get("max_rating"),
            c.get("average_rating"), c.get("rating_percentage"),
            c.get("attainment"), c.get("level_label")
        ]
        for c_idx, v in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=v)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx != 7 else "left")
        r_idx += 1

    auto_fit_columns(ws)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def create_copo_only_excel(copo_data: Dict[str, Any], teacher: Dict[str, Any]) -> io.BytesIO:
    """Generates CO-PO mapping & PO attainment Excel report."""
    wb = Workbook()
    ws_po = wb.active
    ws_po.title = "PO Attainment"
    ws_po.views.sheetView[0].showGridLines = True

    ws_po["A1"] = "D. Y. PATIL PRATISHTHAN'S D. Y. PATIL COLLEGE OF ENGINEERING, AKURDI, PUNE - 411044"
    ws_po["A1"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    ws_po["A2"] = "Approved by AICTE, Recognized by DTE, Affiliated to SPPU, Accredited by NAAC with 'A' Grade, NBA Accredited, Autonomous AY 2024-25"
    ws_po["A2"].font = Font(name="Calibri", size=9, italic=True, color="475569")

    ws_po["A4"] = f"Program Outcome Attainment - {teacher.get('subject', 'Subject')}"
    ws_po["A4"].font = TITLE_FONT

    headers = ["Program Outcome", "Mapping Level Sum", "Weighted CO Sum", "PO Attainment", "Attainment Level"]
    for c_idx, h in enumerate(headers, start=1):
        ws_po.cell(row=6, column=c_idx, value=h)
    style_header_row(ws_po, 6, len(headers))

    r_idx = 7
    for p in copo_data.get("po_results", []):
        att = p.get("attainment")
        att_disp = att if att is not None else "N/A"
        row_vals = [p.get("po"), p.get("mapping_sum"), p.get("weighted_sum"), att_disp, p.get("level_label")]
        for c_idx, v in enumerate(row_vals, start=1):
            cell = ws_po.cell(row=r_idx, column=c_idx, value=v)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx != 5 else "left")
        r_idx += 1

    auto_fit_columns(ws_po)

    # Mapping Matrix sheet
    ws_map = wb.create_sheet(title="CO-PO Mapping Matrix")
    ws_map.views.sheetView[0].showGridLines = True
    po_list = copo_data.get("po_list", [])
    co_list = copo_data.get("co_list", [])
    matrix = copo_data.get("matrix", {})

    m_hdrs = ["Course Outcome"] + po_list
    for c_idx, h in enumerate(m_hdrs, start=1):
        ws_map.cell(row=1, column=c_idx, value=h)
    style_header_row(ws_map, 1, len(m_hdrs))

    r_idx = 2
    for co in co_list:
        ws_map.cell(row=r_idx, column=1, value=co).font = BOLD_FONT
        ws_map.cell(row=r_idx, column=1).border = CELL_BORDER
        for c_idx, po in enumerate(po_list, start=2):
            val = matrix.get(co, {}).get(po, 0)
            cell = ws_map.cell(row=r_idx, column=c_idx, value=val if val != 0 else "-")
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center")
        r_idx += 1

    auto_fit_columns(ws_map)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out


def create_students_only_excel(direct_data: Dict[str, Any], teacher: Dict[str, Any]) -> io.BytesIO:
    """Generates dedicated student performance & marks Excel report."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Performance"
    ws.views.sheetView[0].showGridLines = True

    direct_data = direct_data or {}
    teacher = teacher or {}

    ws["A1"] = "D. Y. PATIL PRATISHTHAN'S D. Y. PATIL COLLEGE OF ENGINEERING, AKURDI, PUNE - 411044"
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    ws["A2"] = "(Approved by AICTE, Recognized by DTE, Affiliated to SPPU, Accredited by NAAC with 'A' Grade, NBA Accredited, Autonomous AY 2024-25)"
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="475569")

    ws["A4"] = f"Student Performance & Attainment Report - {teacher.get('subject', 'Course')}"
    ws["A4"].font = TITLE_FONT

    meta_info = f"Faculty: {teacher.get('faculty_name', 'N/A')} | Class: {teacher.get('class_name', 'N/A')} ({teacher.get('division', 'A')}) | Academic Year: {teacher.get('academic_year', 'N/A')}"
    ws["A5"] = meta_info
    ws["A5"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

    active_cols = direct_data.get("active_columns", [])
    col_names = [c["name"] for c in active_cols]

    headers = ["Roll No", "Student Name"] + col_names + ["Total Obtained", "Max Total", "Percentage (%)", "Passed COs", "Status"]
    for c_idx, h in enumerate(headers, start=1):
        ws.cell(row=7, column=c_idx, value=h)
    style_header_row(ws, 7, len(headers))

    r_idx = 8
    for s in direct_data.get("students", []):
        status_label = "All Passed" if s.get("passed_all") else "Needs Attention"
        row_vals = [s.get("roll_no"), s.get("name")]
        for col_name in col_names:
            row_vals.append(s.get("marks", {}).get(col_name, ""))
        row_vals.extend([
            s.get("total_obtained"),
            s.get("total_max"),
            s.get("percentage"),
            ", ".join(s.get("passed_cos", [])),
            status_label
        ])

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center" if c_idx not in [2, len(row_vals) - 1] else "left")
            if r_idx % 2 == 1:
                cell.fill = ZEBRA_FILL
        r_idx += 1

    auto_fit_columns(ws)

    # Sheet 2: Performance Distribution
    ws_dist = wb.create_sheet(title="Score Distribution")
    ws_dist.views.sheetView[0].showGridLines = True
    ws_dist["A1"] = "Student Score Range Breakdown"
    ws_dist["A1"].font = SECTION_FONT

    dist_headers = ["Score Range", "Student Count", "Percentage of Class (%)"]
    for c_idx, h in enumerate(dist_headers, start=1):
        ws_dist.cell(row=3, column=c_idx, value=h)
    style_header_row(ws_dist, 3, len(dist_headers), fill=SUBHEADER_FILL)

    score_ranges = direct_data.get("score_ranges", {})
    total_stud = direct_data.get("total_students", 1) or 1
    r_idx = 4
    for rng, cnt in score_ranges.items():
        pct = round((cnt / total_stud) * 100.0, 1)
        row_v = [rng, cnt, f"{pct}%"]
        for c_idx, v in enumerate(row_v, start=1):
            cell = ws_dist.cell(row=r_idx, column=c_idx, value=v)
            cell.font = REGULAR_FONT
            cell.border = CELL_BORDER
            cell.alignment = Alignment(horizontal="center")
        r_idx += 1

    auto_fit_columns(ws_dist)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out
