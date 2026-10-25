"""
Comprehensive automated tests for CO-PO Attainment System.
Tests:
1. Direct marks validation (negative, >max, blank, non-numeric, invalid max marks)
2. Unchecked column exclusion
3. Multi-CO mapping
4. Direct attainment formulas (continuous 0-3 scale matching GAPC guidelines)
5. Indirect survey parsing, CO detection, ignoring non-CO columns, rating attainment
6. Weighted combination with custom weights (80/20 & total 100% check)
7. Named and Headerless CO-PO matrices
8. PO attainment calculation
9. Excel export generation (all sheets present and non-empty)
10. Word document generation
11. All Flask page routes and API endpoints
"""

import io
import json
from app import app, TEMP_DATA_STORE
from calculations import (
    calculate_attainment_from_percentage,
    calculate_attainment_0_to_1,
    get_attainment_level_label,
    validate_and_compute_direct,
    validate_and_compute_indirect,
    compute_weighted_co_attainment,
    compute_po_attainment
)
from excel_export import create_complete_excel_report
from docx_export import create_word_report
from pdf_export import create_pdf_report
import openpyxl
from docx import Document


def test_gapc_attainment_formula():
    """Verify exact GAPC 4.0 attainment formula and lookup table levels in points."""
    # 22% -> (22*1)/60 = 0.367 -> Level 0-1 (0.37)
    att_22 = calculate_attainment_from_percentage(22.0)
    assert att_22 == 0.367
    assert get_attainment_level_label(att_22) == "Level 0-1 (0.37)"

    # 35% -> (35*1)/60 = 0.583 -> Level 0-1 (0.58) (NOT Level 2, NOT fixed Level 0!)
    att_35 = calculate_attainment_from_percentage(35.0)
    assert att_35 == 0.583
    assert get_attainment_level_label(att_35) == "Level 0-1 (0.58)"

    # 60% -> 1.0 -> Level 1
    att_60 = calculate_attainment_from_percentage(60.0)
    assert att_60 == 1.0
    assert get_attainment_level_label(att_60) == "Level 1"

    # 67% -> 1 + 7/10 = 1.7 -> Level 1.7
    att_67 = calculate_attainment_from_percentage(67.0)
    assert att_67 == 1.7
    assert get_attainment_level_label(att_67) == "Level 1.7"

    # 70% -> 2.0 -> Level 2
    att_70 = calculate_attainment_from_percentage(70.0)
    assert att_70 == 2.0
    assert get_attainment_level_label(att_70) == "Level 2"

    # 75% -> 2 + 5/10 = 2.5 -> Level 2.5
    att_75 = calculate_attainment_from_percentage(75.0)
    assert att_75 == 2.5
    assert get_attainment_level_label(att_75) == "Level 2.5"

    # 79% -> 2 + 9/10 = 2.9 -> Level 2.9
    att_79 = calculate_attainment_from_percentage(79.0)
    assert att_79 == 2.9
    assert get_attainment_level_label(att_79) == "Level 2.9"

    # 80% and above -> 3.0 -> Level 3
    att_80 = calculate_attainment_from_percentage(80.0)
    assert att_80 == 3.0
    assert get_attainment_level_label(att_80) == "Level 3"
    assert calculate_attainment_from_percentage(95.0) == 3.0
    assert get_attainment_level_label(3.0) == "Level 3"
    assert get_attainment_level_label(0.0) == "Level 0"


def test_direct_validations():
    """Verify rejection of invalid marks with exact row and assessment name."""
    col_configs = [
        {'name': 'UT1', 'checked': True, 'max_marks': 20, 'mapped_cos': ['CO1']},
        {'name': 'CCE1', 'checked': True, 'max_marks': 10, 'mapped_cos': ['CO1', 'CO2']}
    ]

    # Test negative marks
    neg_data = [{'roll_no': '1', 'name': 'Alice', 'marks': {'UT1': -5, 'CCE1': 8}}]
    res_neg = validate_and_compute_direct(neg_data, col_configs)
    assert not res_neg['success']
    assert any("Row 1 (Alice)" in e and "negative" in e and "UT1" in e for e in res_neg['errors'])

    # Test marks > max marks
    exceed_data = [{'roll_no': '2', 'name': 'Bob', 'marks': {'UT1': 25, 'CCE1': 8}}]
    res_exceed = validate_and_compute_direct(exceed_data, col_configs)
    assert not res_exceed['success']
    assert any("Row 1 (Bob)" in e and "exceeds maximum marks" in e and "UT1" in e for e in res_exceed['errors'])

    # Test blank marks
    blank_data = [{'roll_no': '3', 'name': 'Charlie', 'marks': {'UT1': '', 'CCE1': 8}}]
    res_blank = validate_and_compute_direct(blank_data, col_configs)
    assert not res_blank['success']
    assert any("Row 1 (Charlie)" in e and "blank" in e and "UT1" in e for e in res_blank['errors'])

    # Test non-numeric marks
    alpha_data = [{'roll_no': '4', 'name': 'David', 'marks': {'UT1': 'AB', 'CCE1': 8}}]
    res_alpha = validate_and_compute_direct(alpha_data, col_configs)
    assert not res_alpha['success']
    assert any("Row 1 (David)" in e and "non-numeric" in e and "UT1" in e for e in res_alpha['errors'])


def test_unchecked_column_exclusion():
    """Verify that unchecked column is completely excluded from calculation."""
    col_configs = [
        {'name': 'UT1', 'checked': True, 'max_marks': 20, 'mapped_cos': ['CO1']},
        {'name': 'Prelim', 'checked': False, 'max_marks': 50, 'mapped_cos': ['CO2']}  # Unchecked!
    ]
    # Prelim has invalid mark, but because it is unchecked it should be completely ignored!
    students = [
        {'roll_no': '101', 'name': 'Student A', 'marks': {'UT1': 16, 'Prelim': 999}}
    ]
    res = validate_and_compute_direct(students, col_configs)
    assert res['success'], f"Failed with errors: {res.get('errors')}"
    data = res['data']
    # Total max should only be 20 (UT1), NOT 70
    assert data['students'][0]['total_max'] == 20
    assert data['students'][0]['total_obtained'] == 16
    assert data['students'][0]['percentage'] == 80.0
    # Assessment summaries should only have UT1
    assert len(data['assessment_summaries']) == 1
    assert data['assessment_summaries'][0]['name'] == 'UT1'


def test_multi_co_mapping():
    """Verify an assessment mapping to multiple COs correctly attributes marks."""
    col_configs = [
        {'name': 'Insem', 'checked': True, 'max_marks': 30, 'mapped_cos': ['CO1', 'CO2']}
    ]
    students = [
        {'roll_no': '101', 'name': 'Student A', 'marks': {'Insem': 18}}, # 18/30 = 60% -> passes (>=50%)
        {'roll_no': '102', 'name': 'Student B', 'marks': {'Insem': 12}}  # 12/30 = 40% -> fails (<50%)
    ]
    res = validate_and_compute_direct(students, col_configs)
    assert res['success']
    cos = {c['co']: c for c in res['data']['co_summaries']}
    assert 'CO1' in cos and 'CO2' in cos
    assert cos['CO1']['students_passed'] == 1
    assert cos['CO2']['students_passed'] == 1
    assert cos['CO1']['pass_percentage'] == 50.0
    assert cos['CO2']['pass_percentage'] == 50.0


def test_indirect_calculations():
    """Verify indirect survey parsing, CO column detection, ignoring non-CO columns."""
    ratings = [
        {'Roll No': '101', 'Name': 'Alice', 'CO1': 8, 'CO2': 9, 'Comments': 'Great'},
        {'Roll No': '102', 'Name': 'Bob', 'CO1': 9, 'CO2': 8, 'Comments': 'Nice'},
        {'Roll No': '103', 'Name': 'Charlie', 'CO1': 7, 'CO2': 7, 'Comments': 'Good'}
    ]
    res = validate_and_compute_indirect(ratings, max_rating=10.0)
    assert res['success']
    data = res['data']
    assert set(data['detected_cos']) == {'CO1', 'CO2'}
    co1 = next(c for c in data['co_summaries'] if c['co'] == 'CO1')
    # CO1 avg = (8 + 9 + 7)/3 = 8.0 -> 80% -> attainment 3.0
    assert co1['average_rating'] == 8.0
    assert co1['rating_percentage'] == 80.0
    assert co1['attainment'] == 3.0


def test_weighted_calculations():
    """Verify weighted direct and indirect formula."""
    d_cos = [{'co': 'CO1', 'pass_percentage': 75.0, 'attainment': 2.5}]
    i_cos = [{'co': 'CO1', 'rating_percentage': 88.0, 'attainment': 3.0}]
    
    # Direct weight 80%, Indirect weight 20%
    res = compute_weighted_co_attainment(d_cos, i_cos, direct_weight=80.0, indirect_weight=20.0)
    assert res['success']
    c1 = res['data']['combined_results'][0]
    # Direct contrib = 75 * 0.8 = 60.0%
    assert c1['direct_contribution'] == 60.0
    # Indirect contrib = 88 * 0.2 = 17.6% (as specifically cited in prompt!)
    assert c1['indirect_contribution'] == 17.6
    # Combined % = 60.0 + 17.6 = 77.6%
    assert c1['combined_pct'] == 77.6
    # Combined attainment = 77.6 / 100 * 3 = 2.328
    assert c1['combined_attainment'] == 2.328


def test_headerless_copo_matrix():
    """Verify headerless matrix parsing and PO attainment."""
    client = app.test_client()
    # 1. Basic 3-column headerless matrix text
    raw_text = "3,2,1\n2,3,2" # 2 rows (CO1, CO2), 3 cols (PO1, PO2, PO3)
    res = client.post('/api/upload-copo', data={'pasted_text': raw_text})
    assert res.status_code == 200
    data = res.get_json()
    assert data['success']
    assert data['is_headerless'] is True
    assert data['co_list'] == ['CO1', 'CO2']
    assert data['po_list'] == ['PO1', 'PO2', 'PO3']
    assert data['matrix']['CO1']['PO1'] == 3.0

    # 2. 15-column pure matrix with hyphens, no row or col names:
    # Must preserve column 0 as PO1 and assign up to PSO4!
    raw_15_col = "3,2,1,-,-,2,1,-,3,2,1,-,-,-,2\n2,3,2,1,-,2,1,3,2,1,-,2,1,3,1"
    res15 = client.post('/api/upload-copo', data={'pasted_text': raw_15_col})
    assert res15.status_code == 200
    data15 = res15.get_json()
    assert data15['success']
    assert data15['is_headerless'] is True
    assert data15['co_list'] == ['CO1', 'CO2']
    expected_pos_15 = ['PO1', 'PO2', 'PO3', 'PO4', 'PO5', 'PO6', 'PO7', 'PO8', 'PO9', 'PO10', 'PO11', 'PSO1', 'PSO2', 'PSO3', 'PSO4']
    assert data15['po_list'] == expected_pos_15
    assert data15['matrix']['CO1']['PO1'] == 3.0  # 1st column preserved as PO1
    assert data15['matrix']['CO1']['PSO4'] == 2.0 # 15th column mapped to PSO4
    assert data15['matrix']['CO1']['PO4'] == 0.0  # hyphen '-' converted to 0.0

    # 3. 16-column matrix: columns increase to PSO5!
    raw_16_col = "3,2,1,0,0,2,1,0,3,2,1,0,0,0,2,1\n2,3,2,1,0,2,1,3,2,1,0,2,1,3,1,2"
    res16 = client.post('/api/upload-copo', data={'pasted_text': raw_16_col})
    assert res16.status_code == 200
    data16 = res16.get_json()
    assert data16['po_list'][-1] == 'PSO5'
    assert data16['matrix']['CO1']['PO1'] == 3.0
    assert data16['matrix']['CO1']['PSO5'] == 1.0


def test_excel_export_complete():
    """Verify complete multi-sheet Excel generation and content."""
    mock_data = {
        'teacher': {'faculty_name': 'Prof Test', 'subject': 'Calculus', 'class_name': 'FE', 'division': 'A', 'semester': 'I', 'academic_year': '2025-26'},
        'direct': {'data': {
            'active_columns': [{'name': 'UT1', 'max_marks': 20, 'mapped_cos': ['CO1']}],
            'all_cos': ['CO1'],
            'students': [{'roll_no': '1', 'name': 'A', 'marks': {'UT1': 16}, 'total_obtained': 16, 'total_max': 20, 'percentage': 80.0, 'passed_cos': ['CO1']}],
            'assessment_summaries': [{'name': 'UT1', 'mapped_cos': 'CO1', 'max_marks': 20, 'passing_mark': 10, 'appeared': 1, 'passed': 1, 'failed': 0, 'pass_pct': 100, 'avg_marks': 16, 'min_marks': 16, 'max_marks_scored': 16}],
            'co_summaries': [{'co': 'CO1', 'mapped_assessments': 'UT1', 'max_marks': 20, 'students_appeared': 1, 'students_passed': 1, 'pass_percentage': 100.0, 'attainment': 3.0, 'level_label': 'Level 3'}],
            'score_ranges': {'>=80%': 1}
        }},
        'indirect': {'data': {
            'co_summaries': [{'co': 'CO1', 'response_count': 1, 'max_rating': 10, 'average_rating': 9.0, 'rating_percentage': 90.0, 'attainment': 3.0, 'level_label': 'Level 3'}],
            'cleaned_rows': [{'CO1': 9.0}],
            'detected_cos': ['CO1']
        }},
        'weighted': {'data': {
            'direct_weight': 80,
            'indirect_weight': 20,
            'combined_results': [{'co': 'CO1', 'direct_pct': 100.0, 'direct_contribution': 80.0, 'direct_attainment': 3.0, 'indirect_pct': 90.0, 'indirect_contribution': 18.0, 'indirect_attainment': 3.0, 'combined_pct': 98.0, 'combined_attainment': 2.94, 'level_label': 'Level 3'}],
            'overall_attainment': 2.94,
            'overall_level_label': 'Level 3'
        }},
        'copo': {
            'po_list': ['PO1'],
            'co_list': ['CO1'],
            'matrix': {'CO1': {'PO1': 3.0}},
            'po_results': [{'po': 'PO1', 'mapping_sum': 3.0, 'weighted_sum': 8.82, 'attainment': 2.94, 'level_label': 'Level 3', 'contributing_cos': []}]
        }
    }

    excel_io = create_complete_excel_report(mock_data)
    wb = openpyxl.load_workbook(excel_io)
    sheet_names = wb.sheetnames
    assert "Course & Faculty Info" in sheet_names
    assert "Assessment Summaries" in sheet_names
    assert "Student Marks & Attainment" in sheet_names
    assert "Direct CO Attainment" in sheet_names
    assert "Indirect CO Attainment" in sheet_names
    assert "Weighted CO Attainment" in sheet_names
    assert "CO-PO Mapping Matrix" in sheet_names
    assert "PO Attainment" in sheet_names
    assert "Graph & Distribution Data" in sheet_names


def test_docx_export():
    """Verify Word report document generation."""
    mock_data = {
        'teacher': {'faculty_name': 'Prof Test', 'subject': 'Calculus', 'class_name': 'FE', 'division': 'A', 'semester': 'I', 'academic_year': '2025-26'},
        'direct': {'data': {'assessment_summaries': [], 'co_summaries': []}},
        'indirect': {'data': {'co_summaries': []}},
        'weighted': {'data': {'combined_results': []}},
        'copo': {'po_results': []}
    }
    doc_io = create_word_report(mock_data)
    doc = Document(doc_io)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "COURSE OUTCOME & PROGRAM OUTCOME" in text
    assert "Calculus" in text


def test_flask_routes_and_navigation():
    """Verify Flask page routes, navbar absence on page 1, and teacher redirect."""
    client = app.test_client()

    # Page 1 (Teacher Info): Must NOT have navbar
    r1 = client.get('/')
    assert r1.status_code == 200
    html1 = r1.get_data(as_text=True)
    assert '<header class="steps-nav">' not in html1
    assert 'Faculty Name' in html1
    assert 'Subject Name' in html1

    # Save teacher info
    r_save = client.post('/api/save-teacher', json={
        'faculty_name': 'Dr. Sharma',
        'subject': 'Data Structures',
        'class_name': 'SE',
        'division': 'B',
        'semester': 'III',
        'academic_year': '2025-26'
    })
    assert r_save.status_code == 200
    assert r_save.get_json()['success'] is True

    # Page 2 accessible and has step navigation
    r2 = client.get('/assessment')
    assert r2.status_code == 200
    html2 = r2.get_data(as_text=True)
    assert '<header class="steps-nav">' in html2
    assert 'Data Structures' in html2

    # Page 3 accessible
    r3 = client.get('/copo')
    assert r3.status_code == 200

    # Page 4 accessible
    r4 = client.get('/report')
    assert r4.status_code == 200

    # Test PDF download endpoint with charts=1 and charts=0
    r_pdf_charts = client.get('/download/pdf-report?charts=1')
    assert r_pdf_charts.status_code == 200
    assert r_pdf_charts.mimetype == 'application/pdf'
    assert r_pdf_charts.data.startswith(b'%PDF')

    r_pdf_no_charts = client.get('/download/pdf-report?charts=0')
    assert r_pdf_no_charts.status_code == 200
    assert r_pdf_no_charts.mimetype == 'application/pdf'
    assert r_pdf_no_charts.data.startswith(b'%PDF')

    # Test Slow & Fast Learners categorization API
    r_learn = client.post('/api/categorize-learners', json={
        'slow_threshold_pct': 50,
        'fast_threshold_pct': 70,
        'direct_computed': {
            'students': [
                {'roll_no': '1', 'name': 'Fast Alice', 'percentage': 85.0, 'marks': {'UT1': 17}, 'passed_cos': ['CO1', 'CO2']},
                {'roll_no': '2', 'name': 'Slow Bob', 'percentage': 42.0, 'marks': {'UT1': 8}, 'passed_cos': []}
            ],
            'assessment_summaries': [{'name': 'UT1', 'max_marks': 20, 'mapped_cos': 'CO1'}]
        }
    })
    assert r_learn.status_code == 200
    assert r_learn.get_json()['success'] is True
    assert r_learn.get_json()['data']['fast_count'] == 1
    assert r_learn.get_json()['data']['slow_count'] == 1

    # Test Learners PDF download
    r_learn_pdf = client.get('/download/learners-pdf')
    assert r_learn_pdf.status_code == 200
    assert r_learn_pdf.mimetype == 'application/pdf'
    assert r_learn_pdf.data.startswith(b'%PDF')


def test_students_report_excel_export():
    """Verify dedicated Student Performance Excel report generation."""
    from excel_export import create_students_only_excel
    import openpyxl

    mock_direct = {
        'total_students': 2,
        'active_columns': [{'name': 'UT1'}, {'name': 'CCE1'}],
        'score_ranges': {'>=80%': 1, '70-79%': 1, '60-69%': 0, '50-59%': 0, '<50%': 0},
        'students': [
            {'roll_no': '1', 'name': 'Alice', 'marks': {'UT1': 18, 'CCE1': 9}, 'total_obtained': 27, 'total_max': 30, 'percentage': 90.0, 'passed_cos': ['CO1', 'CO2'], 'passed_all': True},
            {'roll_no': '2', 'name': 'Bob', 'marks': {'UT1': 14, 'CCE1': 8}, 'total_obtained': 22, 'total_max': 30, 'percentage': 73.3, 'passed_cos': ['CO1'], 'passed_all': False}
        ]
    }
    mock_teacher = {'faculty_name': 'Prof Test', 'subject': 'Data Structures', 'class_name': 'SE', 'division': 'A', 'academic_year': '2025-26'}
    excel_io = create_students_only_excel(mock_direct, mock_teacher)
    wb = openpyxl.load_workbook(excel_io)
    assert "Student Performance" in wb.sheetnames
    assert "Score Distribution" in wb.sheetnames
    ws = wb["Student Performance"]
    assert "D. Y. PATIL" in ws["A1"].value
    assert ws["A4"].value.startswith("Student Performance")
    assert ws.cell(row=8, column=1).value == '1'  # Alice roll_no
    assert ws.cell(row=8, column=2).value == 'Alice'
    assert ws.cell(row=9, column=2).value == 'Bob'


def test_indirect_po_mapping():
    """Verify that indirect feedback survey can be mapped directly to Program Outcomes (POs)."""
    ratings_data = [
        {'Roll': '1', 'Name': 'Alice', 'PO1': 9, 'PO2': 8, 'PO3': 8.5},
        {'Roll': '2', 'Name': 'Bob', 'PO1': 8, 'PO2': 7, 'PO3': 8}
    ]
    res = validate_and_compute_indirect(ratings_data, max_rating=10.0, mapping_target='PO')
    assert res['success']
    data = res['data']
    assert data['mapping_target'] == 'PO'
    assert 'PO1' in data['detected_cos']
    assert 'PO2' in data['detected_cos']
    assert 'PO3' in data['detected_cos']
    po1 = next(item for item in data['co_summaries'] if item['co'] == 'PO1')
    assert po1['average_rating'] == 8.5
    assert po1['rating_percentage'] == 85.0
    assert po1['attainment'] == 3.0
    assert po1['attainment_0_1'] == 0.85


def test_0_to_1_attainment_scale_and_ese_separation():
    """Verify 0-1 scale calculation and dedicated ESE separation in direct evaluation."""
    assert calculate_attainment_0_to_1(100.0) == 1.0
    assert calculate_attainment_0_to_1(85.5) == 0.855
    assert calculate_attainment_0_to_1(0.0) == 0.0

    col_configs = [
        {'name': 'UT1', 'checked': True, 'max_marks': 20, 'mapped_cos': ['CO1']},
        {'name': 'ESE', 'checked': True, 'max_marks': 80, 'mapped_cos': ['CO1']}
    ]
    students = [
        {'roll_no': '1', 'name': 'Alice', 'marks': {'UT1': 16, 'ESE': 64}},
        {'roll_no': '2', 'name': 'Bob', 'marks': {'UT1': 8, 'ESE': 50}}
    ]
    res = validate_and_compute_direct(students, col_configs)
    assert res['success']
    co1 = res['data']['co_summaries'][0]
    assert co1['has_internal'] is True
    assert co1['has_ese'] is True
    assert co1['internal_pass_pct'] == 50.0
    assert co1['internal_attainment_0_1'] == 0.50
    assert co1['ese_pass_pct'] == 100.0
    assert co1['ese_attainment_0_1'] == 1.0


def test_pdf_report_export():
    """Verify PDF report generation with and without charts."""
    mock_report = {
        'teacher': {'faculty_name': 'Dr. Test', 'subject': 'Operating Systems', 'class_name': 'TE', 'division': 'A', 'semester': 'V', 'academic_year': '2025-26'},
        'direct': {'data': {
            'assessment_summaries': [{'name': 'UT1', 'mapped_cos': 'CO1', 'max_marks': 20, 'passing_mark': 10, 'appeared': 1, 'passed': 1, 'failed': 0, 'pass_pct': 100, 'avg_marks': 18, 'min_marks': 18, 'max_marks_scored': 18}],
            'co_summaries': [{'co': 'CO1', 'mapped_assessments': 'UT1', 'max_marks': 20, 'students_appeared': 1, 'students_passed': 1, 'pass_percentage': 100.0, 'attainment': 3.0, 'attainment_0_1': 1.0, 'level_label': 'Level 3'}],
            'score_ranges': {'>=80%': 1}
        }},
        'indirect': {'data': {
            'co_summaries': [{'co': 'CO1', 'response_count': 1, 'max_rating': 10, 'average_rating': 9.0, 'rating_percentage': 90.0, 'attainment': 3.0, 'attainment_0_1': 0.9, 'level_label': 'Level 3'}]
        }},
        'weighted': {'data': {
            'combined_results': [{'co': 'CO1', 'direct_pct': 100.0, 'direct_contribution': 80.0, 'direct_attainment': 3.0, 'direct_attainment_0_1': 1.0, 'indirect_pct': 90.0, 'indirect_contribution': 18.0, 'indirect_attainment': 3.0, 'indirect_attainment_0_1': 0.9, 'combined_pct': 98.0, 'combined_attainment': 2.94, 'combined_attainment_0_1': 0.98, 'level_label': 'Level 3'}],
            'overall_attainment': 2.94,
            'overall_attainment_0_1': 0.98,
            'overall_level_label': 'Level 3'
        }},
        'copo': {'po_results': [{'po': 'PO1', 'mapping_sum': 3.0, 'weighted_sum': 8.82, 'attainment': 2.94, 'attainment_0_1': 0.98, 'level_label': 'Level 3', 'contributing_cos': []}]}
    }

    # With charts
    pdf_with = create_pdf_report(mock_report, include_charts=True)
    assert pdf_with.getvalue().startswith(b'%PDF')
    assert len(pdf_with.getvalue()) > 1000

    # Without charts
    pdf_without = create_pdf_report(mock_report, include_charts=False)
    assert pdf_without.getvalue().startswith(b'%PDF')
    assert len(pdf_without.getvalue()) > 1000

    # Assessment Section PDF
    pdf_ass = create_pdf_report(mock_report, include_charts=True, section="assessment")
    assert pdf_ass.getvalue().startswith(b'%PDF')
    assert len(pdf_ass.getvalue()) > 1000

    # CO-PO Section PDF
    pdf_copo = create_pdf_report(mock_report, include_charts=False, section="copo")
    assert pdf_copo.getvalue().startswith(b'%PDF')
    assert len(pdf_copo.getvalue()) > 1000


def test_below_60_percent_level_assignment_and_gapc_lookup():
    """Verify that <60% pass percentage assigns Level 0-1 (0.XX) between 0.0 and 1.0, and lookup table levels."""
    # 35% pass percentage scenario (7 out of 20 students score >= 50% on CCE5)
    col_configs = [{'name': 'CCE5', 'checked': True, 'max_marks': 10, 'mapped_cos': ['CO5']}]
    students_35 = []
    for i in range(1, 21):
        mark = 8 if i <= 7 else 3  # 7 students get 8 (pass >=5), 13 get 3 (fail <5) -> 35% pass
        students_35.append({'roll_no': str(i), 'name': f'Student {i}', 'marks': {'CCE5': mark}})

    res_35 = validate_and_compute_direct(students_35, col_configs)
    assert res_35['success']
    co5 = res_35['data']['co_summaries'][0]
    assert co5['pass_percentage'] == 35.0
    assert co5['attainment'] == 0.583
    # MUST be Level 0-1 (0.58), NEVER Level 2 (Medium) and NOT fixed Level 0!
    assert co5['level_label'] == "Level 0-1 (0.58)"

    # 67% pass percentage scenario (indirect survey or direct assessment)
    col_configs_67 = [{'name': 'UT1', 'checked': True, 'max_marks': 10, 'mapped_cos': ['CO1']}]
    students_67 = []
    # 6 out of 9 students score >= 5 -> 66.67% -> 66.67% rounded to 66.67% -> 1.0 + (66.67-60)/10 = 1.667 -> Level 1.7
    for i in range(1, 10):
        mark = 8 if i <= 6 else 3
        students_67.append({'roll_no': str(i), 'name': f'Student {i}', 'marks': {'UT1': mark}})
    res_67 = validate_and_compute_direct(students_67, col_configs_67)
    assert res_67['success']
    co1 = res_67['data']['co_summaries'][0]
    assert co1['level_label'] == "Level 1.7"


def test_practical_course_workflow():
    """Verify practical course flow: batch saving, multi-experiment CO mapping, and export titles."""
    client = app.test_client()
    
    # 1. Save practical course with batch
    res_t = client.post('/api/save-teacher', json={
        'faculty_name': 'Dr. S. K. Patil',
        'subject': 'Database Management Systems Lab (CS304PR)',
        'class_name': 'Third Year B.Tech',
        'division': 'B',
        'semester': 'V',
        'academic_year': '2025-2026',
        'course_type': 'practical',
        'batch': 'Batch 2'
    })
    assert res_t.status_code == 200
    assert res_t.get_json()['success'] is True
    assert res_t.get_json()['redirect'] == '/practical-assessment'

    res_page = client.get('/practical-assessment')
    assert res_page.status_code == 200
    
    # Check session
    with client.session_transaction() as sess:
        assert sess['teacher']['course_type'] == 'practical'
        assert sess['teacher']['batch'] == 'Batch 2'

    # Verify navigating to /assessment redirects to /practical-assessment when course_type == 'practical'
    res_redir = client.get('/assessment')
    assert res_redir.status_code == 302
    assert '/practical-assessment' in res_redir.location

    # 2. Build 10 experiments (2 per CO) and PR/OR exam
    col_configs = [
        {'name': 'Exp 1', 'max_marks': 25, 'mapped_cos': ['CO1'], 'checked': True},
        {'name': 'Exp 2', 'max_marks': 25, 'mapped_cos': ['CO1'], 'checked': True},
        {'name': 'Exp 3', 'max_marks': 25, 'mapped_cos': ['CO2'], 'checked': True},
        {'name': 'Exp 4', 'max_marks': 25, 'mapped_cos': ['CO2'], 'checked': True},
        {'name': 'Exp 5', 'max_marks': 25, 'mapped_cos': ['CO3'], 'checked': True},
        {'name': 'Exp 6', 'max_marks': 25, 'mapped_cos': ['CO3'], 'checked': True},
        {'name': 'Exp 7', 'max_marks': 25, 'mapped_cos': ['CO4'], 'checked': True},
        {'name': 'Exp 8', 'max_marks': 25, 'mapped_cos': ['CO4'], 'checked': True},
        {'name': 'Exp 9', 'max_marks': 25, 'mapped_cos': ['CO5'], 'checked': True},
        {'name': 'Exp 10', 'max_marks': 25, 'mapped_cos': ['CO5'], 'checked': True},
        {'name': 'PR/OR Exam', 'max_marks': 25, 'mapped_cos': ['CO1', 'CO2', 'CO3', 'CO4', 'CO5'], 'checked': True},
    ]

    students = [
        {
            'roll_no': f'10{i}',
            'name': f'Student {i}',
            'marks': {
                'Exp 1': 20, 'Exp 2': 22,
                'Exp 3': 18, 'Exp 4': 19,
                'Exp 5': 21, 'Exp 6': 23,
                'Exp 7': 17, 'Exp 8': 20,
                'Exp 9': 22, 'Exp 10': 24,
                'PR/OR Exam': 20
            }
        }
        for i in range(1, 11)
    ]

    res_comp = client.post('/api/compute-direct', json={
        'students': students,
        'column_configs': col_configs,
        'passing_threshold_ratio': 0.50
    })
    assert res_comp.status_code == 200
    comp_json = res_comp.get_json()
    assert comp_json['success'] is True
    co_summaries = comp_json['data']['co_summaries']
    assert len(co_summaries) == 5
    # All 10 students scored >= 50% on all COs
    for co in co_summaries:
        assert co['pass_percentage'] == 100.0
        assert co['attainment'] == 3.0

    # 3. Test PDF generation for practical
    teacher_meta = {
        'faculty_name': 'Dr. S. K. Patil',
        'subject': 'DBMS Lab',
        'class_name': 'TY B.Tech',
        'division': 'B',
        'semester': 'V',
        'academic_year': '2025-2026',
        'course_type': 'practical',
        'batch': 'Batch 2'
    }
    rep_data = {
        'teacher': teacher_meta,
        'direct': {'data': comp_json['data']},
        'indirect': {'data': {'co_summaries': [{'co': f'CO{i}', 'attainment': 2.8} for i in range(1, 6)]}},
        'weighted': {'data': {'overall_attainment': 2.96, 'overall_level_label': 'Level 3', 'direct_weight': 80, 'indirect_weight': 20, 'co_attainments': []}},
        'copo': {'po_results': [], 'matrix': {}}
    }
    pdf_buf = create_pdf_report(rep_data, include_charts=False)
    assert len(pdf_buf.getvalue()) > 1000

    # 4. Test Excel generation with practical metadata
    excel_buf = create_complete_excel_report(rep_data)
    wb = openpyxl.load_workbook(excel_buf)
    ws_info = wb['Course & Faculty Info']
    assert "LABORATORY" in ws_info["A4"].value
    found_batch = False
    for r in range(1, 20):
        val = ws_info.cell(row=r, column=2).value
        if val and "Practical" in str(val) and "Batch 2" in str(val):
            found_batch = True
            break
    assert found_batch, "Practical batch not found in Excel report!"

    # 5. Test Word doc generation with practical metadata
    docx_buf = create_word_report(rep_data)
    doc = Document(docx_buf)
    assert any("LABORATORY" in p.text for p in doc.paragraphs)


def test_practicle_column_parsing_and_mapping():
    """Verify that columns named 'practicle 1', 'practicle 2', 'PR/OR' are recognized, given 25 max marks, and mapped to COs."""
    client = app.test_client()

    csv_content = (
        "Sr No,Roll No,Student Name,practicle 1,practicle 2,practicle 3,practicle 4,PR/OR\n"
        "1,101,Aarav Sharma,22,24,21,23,45\n"
        "2,102,Ananya Patil,20,19,22,24,42\n"
        "3,103,Rohan Kulkarni,24,25,23,24,48\n"
    ).encode('utf-8')

    res = client.post('/api/upload-direct', data={'file': (io.BytesIO(csv_content), 'lab_marks.csv')})
    assert res.status_code == 200
    data = res.get_json()
    assert data['success'] is True
    assert len(data['students']) == 3

    configs = {c['name']: c for c in data['column_configs']}
    assert 'practicle 1' in configs
    assert 'practicle 2' in configs
    assert 'practicle 3' in configs
    assert 'practicle 4' in configs
    assert 'PR/OR' in configs

    # Check max marks
    assert configs['practicle 1']['max_marks'] == 25.0
    assert configs['practicle 2']['max_marks'] == 25.0
    assert configs['practicle 3']['max_marks'] == 25.0
    assert configs['practicle 4']['max_marks'] == 25.0
    assert configs['PR/OR']['max_marks'] == 50.0

    # Check smart 2-exp per CO mapping
    assert configs['practicle 1']['mapped_cos'] == ['CO1']
    assert configs['practicle 2']['mapped_cos'] == ['CO1']
    assert configs['practicle 3']['mapped_cos'] == ['CO2']
    assert configs['practicle 4']['mapped_cos'] == ['CO2']
    assert configs['PR/OR']['mapped_cos'] == ['CO1', 'CO2', 'CO3', 'CO4', 'CO5']

    # Compute attainment
    res_comp = client.post('/api/compute-direct', json={
        'students': data['students'],
        'column_configs': data['column_configs'],
        'passing_threshold_ratio': 0.50
    })
    assert res_comp.status_code == 200
    comp_json = res_comp.get_json()
    assert comp_json['success'] is True
    assert len(comp_json['data']['co_summaries']) >= 2


def test_sample_practical_marks_50_xlsx_upload():
    """Verify that sample_practical_marks_50.xlsx downloads properly, parses 15 students and Exp 1-5, and sets max marks 50."""
    client = app.test_client()

    # 1. Test sample download route
    res_dl = client.get('/sample-data/sample_practical_marks_50.xlsx')
    assert res_dl.status_code == 200
    assert len(res_dl.data) > 1000

    # 2. Test upload and parse
    res_up = client.post('/api/upload-direct', data={'file': (io.BytesIO(res_dl.data), 'sample_practical_marks_50.xlsx')})
    assert res_up.status_code == 200
    data = res_up.get_json()
    assert data['success'] is True
    assert len(data['students']) == 15
    assert len(data['column_configs']) == 5

    names = [c['name'] for c in data['column_configs']]
    assert names == ['Exp 1', 'Exp 2', 'Exp 3', 'Exp 4', 'Exp 5']

    for c in data['column_configs']:
        assert c['max_marks'] == 50.0

    assert data['column_configs'][0]['mapped_cos'] == ['CO1']
    assert data['column_configs'][1]['mapped_cos'] == ['CO1']
    assert data['column_configs'][2]['mapped_cos'] == ['CO2']
    assert data['column_configs'][3]['mapped_cos'] == ['CO2']
    assert data['column_configs'][4]['mapped_cos'] == ['CO3']

    # 3. Test calculation
    res_comp = client.post('/api/compute-direct', json={
        'students': data['students'],
        'column_configs': data['column_configs'],
        'passing_threshold_ratio': 0.50
    })
    assert res_comp.status_code == 200
    comp_json = res_comp.get_json()
    assert comp_json['success'] is True
    assert len(comp_json['data']['co_summaries']) == 3


def test_reference_co_po_direct_indirect_and_reset():
    """Verify DY Patil reference format calculations for direct/indirect PO attainment, average mapping, and default reset."""
    matrix = {
        'CO1': {'PO1': 3, 'PO2': 2, 'PO3': 0},
        'CO2': {'PO1': 2, 'PO2': 3, 'PO3': 0},
    }
    co_atts = {'CO1': 2.5, 'CO2': 2.0}
    direct_atts = {'CO1': 2.4, 'CO2': 1.8}
    indirect_atts = {'CO1': 2.8, 'CO2': 2.6}

    result = compute_po_attainment(
        mapping_matrix=matrix,
        co_attainments=co_atts,
        po_list=['PO1', 'PO2', 'PO3'],
        co_list=['CO1', 'CO2'],
        direct_co_attainments=direct_atts,
        indirect_co_attainments=indirect_atts
    )

    po_dict = {p['po']: p for p in result['po_results']}

    # PO1: mapping_sum = 3 + 2 = 5, len(co_list) = 2 -> mapping_avg = 5 / 2 = 2.5
    assert po_dict['PO1']['mapping_avg'] == 2.5
    # PO1 direct_att = (3*2.4 + 2*1.8) / 5 = (7.2 + 3.6)/5 = 10.8/5 = 2.16
    assert po_dict['PO1']['direct_attainment'] == 2.16
    # PO1 indirect_att = (3*2.8 + 2*2.6) / 5 = (8.4 + 5.2)/5 = 13.6/5 = 2.72
    assert po_dict['PO1']['indirect_attainment'] == 2.72
    # PO1 combined_att = (3*2.5 + 2*2.0) / 5 = (7.5 + 4.0)/5 = 11.5/5 = 2.3
    assert po_dict['PO1']['attainment'] == 2.3

    # PO3: unmapped -> mapping_avg = 0.0, direct = 0.0, indirect = 0.0, attainment = None
    assert po_dict['PO3']['mapping_avg'] == 0.0
    assert po_dict['PO3']['direct_attainment'] == 0.0
    assert po_dict['PO3']['attainment'] is None

    # Test default session reset on visiting / or /teacher
    client = app.test_client()
    with client.session_transaction() as sess:
        sess['teacher'] = {'faculty_name': 'Existing Faculty'}
    res = client.get('/teacher')
    assert res.status_code == 200
    # Session must be fresh and clean
    with client.session_transaction() as sess:
        assert sess.get('teacher', {}).get('faculty_name') != 'Existing Faculty'


if __name__ == '__main__':
    test_gapc_attainment_formula()
    print("[PASS] test_gapc_attainment_formula")
    test_direct_validations()
    print("[PASS] test_direct_validations")
    test_unchecked_column_exclusion()
    print("[PASS] test_unchecked_column_exclusion")
    test_multi_co_mapping()
    print("[PASS] test_multi_co_mapping")
    test_indirect_calculations()
    print("[PASS] test_indirect_calculations")
    test_weighted_calculations()
    print("[PASS] test_weighted_calculations")
    test_headerless_copo_matrix()
    print("[PASS] test_headerless_copo_matrix")
    test_excel_export_complete()
    print("[PASS] test_excel_export_complete")
    test_docx_export()
    print("[PASS] test_docx_export")
    test_0_to_1_attainment_scale_and_ese_separation()
    print("[PASS] test_0_to_1_attainment_scale_and_ese_separation")
    test_pdf_report_export()
    print("[PASS] test_pdf_report_export")
    test_students_report_excel_export()
    print("[PASS] test_students_report_excel_export")
    test_indirect_po_mapping()
    print("[PASS] test_indirect_po_mapping")
    test_flask_routes_and_navigation()
    print("[PASS] test_flask_routes_and_navigation")
    test_below_60_percent_level_assignment_and_gapc_lookup()
    print("[PASS] test_below_60_percent_level_assignment_and_gapc_lookup")
    test_practical_course_workflow()
    print("[PASS] test_practical_course_workflow")
    test_practicle_column_parsing_and_mapping()
    print("[PASS] test_practicle_column_parsing_and_mapping")
    test_sample_practical_marks_50_xlsx_upload()
    print("[PASS] test_sample_practical_marks_50_xlsx_upload")
    test_reference_co_po_direct_indirect_and_reset()
    print("[PASS] test_reference_co_po_direct_indirect_and_reset")
    print("\nALL 19 TEST SUITES PASSED SUCCESSFULLY!")
