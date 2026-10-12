"""
End-to-end integration test verifying full workflow from Page 1 to Page 4 downloads.
"""

from app import app, TEMP_DATA_STORE
import openpyxl
from docx import Document
import io

def run_e2e_test():
    client = app.test_client()

    print("Step 1: Save Teacher & Course info")
    res1 = client.post('/api/save-teacher', json={
        'faculty_name': 'Prof. Aniruddh Joshi',
        'subject': 'Data Structures & Algorithms (CS302)',
        'class_name': 'Third Year B.Tech',
        'division': 'A',
        'semester': 'V',
        'academic_year': '2025-2026'
    })
    assert res1.status_code == 200
    assert res1.get_json()['success'] is True

    print("Step 2: Upload Direct marks (wide format)")
    with open('sample_data/sample_direct_marks_wide.csv', 'rb') as f:
        csv_data = f.read()
    res2 = client.post('/api/upload-direct', data={'file': (io.BytesIO(csv_data), 'marks.csv')})
    assert res2.status_code == 200
    direct_parsed = res2.get_json()
    assert direct_parsed['success'] is True
    assert len(direct_parsed['students']) == 15
    assert len(direct_parsed['column_configs']) == 7 # UT1, CCE1, Insem, UT2, CCE2, Prelim, ESE

    print("Step 2b: Uncheck 'Prelim' and compute Direct Attainment")
    col_configs = direct_parsed['column_configs']
    for c in col_configs:
        if c['name'] == 'Prelim':
            c['checked'] = False # Completely exclude Prelim!
    
    res_comp_direct = client.post('/api/compute-direct', json={
        'students': direct_parsed['students'],
        'column_configs': col_configs,
        'passing_threshold_ratio': 0.50
    })
    assert res_comp_direct.status_code == 200
    comp_direct = res_comp_direct.get_json()
    assert comp_direct['success'] is True
    # Confirm Prelim is not in active columns or summaries
    active_names = [c['name'] for c in comp_direct['data']['active_columns']]
    assert 'Prelim' not in active_names
    print(f"Direct COs evaluated: {comp_direct['data']['all_cos']}")

    print("Step 2c: Upload Indirect survey data")
    with open('sample_data/sample_indirect_feedback.csv', 'rb') as f:
        ind_data = f.read()
    res_ind = client.post('/api/upload-indirect', data={'file': (io.BytesIO(ind_data), 'feedback.csv')})
    assert res_ind.status_code == 200
    ind_parsed = res_ind.get_json()
    assert ind_parsed['success'] is True
    assert 'CO1' in ind_parsed['detected_cos']
    # Non-CO columns like 'Comments' should not be in detected_cos
    assert 'Comments' not in ind_parsed['detected_cos']

    print("Step 2d: Compute Indirect Attainment")
    res_comp_ind = client.post('/api/compute-indirect', json={
        'ratings_data': ind_parsed['ratings_data'],
        'max_rating': 10.0,
        'co_columns_override': ind_parsed['detected_cos']
    })
    assert res_comp_ind.status_code == 200
    comp_ind = res_comp_ind.get_json()
    assert comp_ind['success'] is True

    print("Step 2e: Compute Weighted Attainment (80% Direct / 20% Indirect)")
    res_wt = client.post('/api/compute-weighted', json={'direct_weight': 80.0, 'indirect_weight': 20.0})
    assert res_wt.status_code == 200
    comp_wt = res_wt.get_json()
    assert comp_wt['success'] is True
    print(f"Overall Course Attainment: {comp_wt['data']['overall_attainment']} / 3.00 ({comp_wt['data']['overall_level_label']})")

    print("Step 3: Upload CO-PO Mapping Matrix (Named)")
    with open('sample_data/sample_copo_mapping.csv', 'rb') as f:
        copo_csv = f.read()
    res_copo = client.post('/api/upload-copo', data={'file': (io.BytesIO(copo_csv), 'matrix.csv')})
    assert res_copo.status_code == 200
    copo_parsed = res_copo.get_json()
    assert copo_parsed['success'] is True
    assert 'PO1' in copo_parsed['po_list']

    print("Step 3b: Compute PO Attainment")
    res_comp_po = client.post('/api/compute-copo', json={
        'matrix': copo_parsed['matrix'],
        'po_list': copo_parsed['po_list'],
        'co_list': copo_parsed['co_list']
    })
    assert res_comp_po.status_code == 200
    comp_po = res_comp_po.get_json()
    assert comp_po['success'] is True
    print(f"Overall PO Attainment: {comp_po['data']['overall_po_attainment']} / 3.00 ({comp_po['data']['overall_po_level']})")

    print("Step 4: Verify Final Report State API")
    res_state = client.get('/api/get-full-state')
    assert res_state.status_code == 200
    state = res_state.get_json()
    assert state['teacher']['faculty_name'] == 'Prof. Aniruddh Joshi'
    assert state['weighted']['overall_attainment'] is not None

    print("Step 5: Test Download Complete Excel Report")
    res_dl_xl = client.get('/download/complete-excel')
    assert res_dl_xl.status_code == 200
    assert 'spreadsheetml' in res_dl_xl.mimetype
    wb = openpyxl.load_workbook(io.BytesIO(res_dl_xl.data))
    assert len(wb.sheetnames) >= 9
    print(f"Excel report sheets generated: {wb.sheetnames}")

    print("Step 6: Test Download Word Document Report")
    res_dl_doc = client.get('/download/word-report')
    assert res_dl_doc.status_code == 200
    assert 'wordprocessingml' in res_dl_doc.mimetype
    doc = Document(io.BytesIO(res_dl_doc.data))
    assert len(doc.tables) >= 6
    print(f"Word report generated with {len(doc.tables)} accreditation tables.")

    print("Step 7: Test Individual Reports")
    res_dir = client.get('/download/direct-report')
    assert res_dir.status_code == 200
    res_stud = client.get('/download/students-report')
    assert res_stud.status_code == 200
    assert 'spreadsheetml' in res_stud.mimetype
    res_ind_dl = client.get('/download/indirect-report')
    assert res_ind_dl.status_code == 200
    res_copo_dl = client.get('/download/copo-report')
    assert res_copo_dl.status_code == 200

    print("Step 8: Test Download PDF Document Report (With & Without Charts)")
    res_pdf_with = client.get('/download/pdf-report?charts=1')
    assert res_pdf_with.status_code == 200
    assert res_pdf_with.mimetype == 'application/pdf'
    assert res_pdf_with.data.startswith(b'%PDF')

    res_pdf_without = client.get('/download/pdf-report?charts=0')
    assert res_pdf_without.status_code == 200
    assert res_pdf_without.mimetype == 'application/pdf'
    assert res_pdf_without.data.startswith(b'%PDF')
    print(f"PDF reports generated successfully ({len(res_pdf_with.data)} bytes with charts, {len(res_pdf_without.data)} bytes without charts).")

    print("Step 9: Test Page-Specific PDF Exports (Page 2 Assessment PDF & Page 3 CO-PO PDF)")
    res_ass_pdf = client.get('/download/assessment-pdf?charts=1')
    assert res_ass_pdf.status_code == 200
    assert res_ass_pdf.mimetype == 'application/pdf'
    assert res_ass_pdf.data.startswith(b'%PDF')

    res_copo_pdf = client.get('/download/copo-pdf')
    assert res_copo_pdf.status_code == 200
    assert res_copo_pdf.mimetype == 'application/pdf'
    assert res_copo_pdf.data.startswith(b'%PDF')
    print(f"Page-specific PDFs generated successfully (Assessment: {len(res_ass_pdf.data)} bytes, CO-PO: {len(res_copo_pdf.data)} bytes).")

    print("\nALL END-TO-END WORKFLOW CHECKS PASSED WITH FLYING COLORS!")

if __name__ == '__main__':
    run_e2e_test()
