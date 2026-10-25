"""
Word (DOCX) Export Generator for CO-PO Attainment System.
Generates fully styled, professional documents for university accreditation (NBA/NAAC/GAPC).
"""

import io
from typing import Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex: str):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def style_table(table, col_widths=None):
    """Applies clean borders and formatting to a python-docx table."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        # Prevent row splitting across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if row_idx == 0:
            # Repeat header row on each page
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for c_idx, cell in enumerate(row.cells):
            # Padding
            tcPr = cell._tc.get_or_add_tcPr()
            tcMar = parse_xml(
                f'<w:tcMar {nsdecls("w")}>'
                f'<w:top w:w="120" w:type="dxa"/>'
                f'<w:bottom w:w="120" w:type="dxa"/>'
                f'<w:left w:w="150" w:type="dxa"/>'
                f'<w:right w:w="150" w:type="dxa"/>'
                f'</w:tcMar>'
            )
            tcPr.append(tcMar)
            
            # Subtle cell borders
            borders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(borders)

            if col_widths and c_idx < len(col_widths):
                cell.width = Inches(col_widths[c_idx])


def create_word_report(report_data: Dict[str, Any]) -> io.BytesIO:
    """Creates a complete DOCX report."""
    doc = Document()
    
    # Page setup - standard Letter / A4 with 0.8 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    teacher = report_data.get('teacher') or {}
    direct = (report_data.get('direct') or {}).get('data') or {}
    indirect = (report_data.get('indirect') or {}).get('data') or {}
    weighted = (report_data.get('weighted') or {}).get('data') or {}
    copo = report_data.get('copo') or {}
    po_results = copo.get('po_results') or []
    mapping_matrix = copo.get('matrix') or {}

    # Institutional Accreditation Header
    inst_p = doc.add_paragraph()
    inst_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_inst = inst_p.add_run("D. Y. PATIL PRATISHTHAN'S\nD. Y. PATIL COLLEGE OF ENGINEERING, AKURDI, PUNE - 411044\n")
    run_inst.font.name = "Arial"
    run_inst.font.size = Pt(13)
    run_inst.font.bold = True
    run_inst.font.color.rgb = RGBColor(30, 58, 138)
    
    run_subinst = inst_p.add_run("(Approved by AICTE, Recognized by DTE, Affiliated to SPPU, Accredited by NAAC with 'A' Grade, NBA Accredited, Autonomous AY 2024-25)")
    run_subinst.font.name = "Arial"
    run_subinst.font.size = Pt(8.5)
    run_subinst.font.italic = True
    run_subinst.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc_title = "LABORATORY COURSE OUTCOME & PROGRAM OUTCOME ATTAINMENT REPORT" if teacher.get("course_type") == "practical" else "COURSE OUTCOME & PROGRAM OUTCOME ATTAINMENT REPORT"
    run_t = title_p.add_run(doc_title)
    run_t.font.name = "Arial"
    run_t.font.size = Pt(15)
    run_t.font.bold = True
    run_t.font.color.rgb = RGBColor(30, 41, 59)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    course_type_str = f"Practical (Batch: {teacher.get('batch')})" if teacher.get("course_type") == "practical" and teacher.get("batch") else ("Practical" if teacher.get("course_type") == "practical" else "Theory")
    run_s = sub_p.add_run(f"Subject: {teacher.get('subject', 'N/A')}  |  Type: {course_type_str}  |  Academic Year: {teacher.get('academic_year', 'N/A')}")
    run_s.font.size = Pt(10.5)
    run_s.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph()

    # Section: Course & Faculty Details
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Course & Faculty Details")
    h1_run.font.color.rgb = RGBColor(30, 41, 59)

    meta_info = [
        ("Faculty Name", teacher.get("faculty_name", "N/A")),
        ("Subject Name / Code", teacher.get("subject", "N/A")),
        ("Course Nature & Batch", course_type_str),
        ("Class & Division", f"{teacher.get('class_name', 'N/A')} - Div {teacher.get('division', 'N/A')}"),
        ("Semester", teacher.get("semester", "N/A")),
        ("Weightage Setup", f"Direct: {weighted.get('direct_weight', 80)}%  |  Indirect: {weighted.get('indirect_weight', 20)}%"),
        ("Overall Attainment", f"{weighted.get('overall_attainment', 'N/A')} / 3.00  ({weighted.get('overall_level_label', 'N/A')})")
    ]
    table_meta = doc.add_table(rows=len(meta_info), cols=2)
    for r_idx, (k, v) in enumerate(meta_info):
        r = table_meta.rows[r_idx]
        cell_k, cell_v = r.cells[0], r.cells[1]
        set_cell_background(cell_k, "F1F5F9")
        p_k = cell_k.paragraphs[0]
        r_k = p_k.add_run(k)
        r_k.bold = True
        r_k.font.size = Pt(9.5)
        p_v = cell_v.paragraphs[0]
        r_v = p_v.add_run(str(v))
        r_v.font.size = Pt(9.5)

    style_table(table_meta, [2.5, 4.2])
    doc.add_paragraph()

    # Section 2: Direct Assessments Summary
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. Assessment Tools Setup & Passing Summaries")
    h2_run.font.color.rgb = RGBColor(30, 41, 59)

    assess_list = direct.get("assessment_summaries", [])
    if assess_list:
        table_ass = doc.add_table(rows=len(assess_list) + 1, cols=6)
        hdrs = ["Assessment", "Mapped COs", "Max Marks", "Passing Mark", "Appeared", "Pass %"]
        for c_idx, h in enumerate(hdrs):
            cell = table_ass.rows[0].cells[c_idx]
            set_cell_background(cell, "1E293B")
            p = cell.paragraphs[0]
            run = p.add_run(h)
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r_idx, a in enumerate(assess_list, start=1):
            row = table_ass.rows[r_idx]
            vals = [a.get("name"), a.get("mapped_cos"), str(a.get("max_marks")), str(a.get("passing_mark")), str(a.get("appeared")), f"{a.get('pass_pct')}%"]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                if r_idx % 2 == 0:
                    set_cell_background(cell, "F8FAFC")
                p = cell.paragraphs[0]
                run = p.add_run(str(val))
                run.font.size = Pt(9)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 1 else WD_ALIGN_PARAGRAPH.LEFT
        style_table(table_ass, [1.5, 1.2, 0.9, 1.0, 0.9, 1.1])
    doc.add_paragraph()

    # Section 3: Direct CO Attainment
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("3. Direct Course Outcome (CO) Attainment")
    h3_run.font.color.rgb = RGBColor(30, 41, 59)

    direct_cos = direct.get("co_summaries", [])
    if direct_cos:
        table_dco = doc.add_table(rows=len(direct_cos) + 1, cols=6)
        hdrs = ["CO", "Mapped Tests", "Appeared", "Passed (>=50%)", "Pass %", "Direct Attainment"]
        for c_idx, h in enumerate(hdrs):
            cell = table_dco.rows[0].cells[c_idx]
            set_cell_background(cell, "1E293B")
            p = cell.paragraphs[0]
            run = p.add_run(h)
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r_idx, c in enumerate(direct_cos, start=1):
            row = table_dco.rows[r_idx]
            vals = [c.get("co"), c.get("mapped_assessments"), str(c.get("students_appeared")), str(c.get("students_passed")), f"{c.get('pass_percentage')}%", f"{c.get('attainment')} ({c.get('level_label')})"]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                if r_idx % 2 == 0:
                    set_cell_background(cell, "F8FAFC")
                p = cell.paragraphs[0]
                run = p.add_run(str(val))
                run.font.size = Pt(9)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx not in [1, 5] else WD_ALIGN_PARAGRAPH.LEFT
        style_table(table_dco, [0.8, 1.8, 0.9, 1.1, 0.9, 1.6])
    doc.add_paragraph()

    # Section 4: Indirect Survey Attainment
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run("4. Indirect Assessment (Survey/Feedback)")
    h4_run.font.color.rgb = RGBColor(30, 41, 59)

    ind_cos = indirect.get("co_summaries", [])
    if ind_cos:
        table_ico = doc.add_table(rows=len(ind_cos) + 1, cols=6)
        hdrs = ["CO", "Responses", "Max Scale", "Average Rating", "Rating %", "Indirect Attainment"]
        for c_idx, h in enumerate(hdrs):
            cell = table_ico.rows[0].cells[c_idx]
            set_cell_background(cell, "1E293B")
            p = cell.paragraphs[0]
            run = p.add_run(h)
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r_idx, c in enumerate(ind_cos, start=1):
            row = table_ico.rows[r_idx]
            vals = [c.get("co"), str(c.get("response_count")), str(c.get("max_rating")), str(c.get("average_rating")), f"{c.get('rating_percentage')}%", f"{c.get('attainment')} ({c.get('level_label')})"]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                if r_idx % 2 == 0:
                    set_cell_background(cell, "F8FAFC")
                p = cell.paragraphs[0]
                run = p.add_run(str(val))
                run.font.size = Pt(9)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != 5 else WD_ALIGN_PARAGRAPH.LEFT
        style_table(table_ico, [0.8, 1.0, 1.0, 1.1, 1.0, 1.7])
    doc.add_paragraph()

    # Section 5: Weighted Combined Attainment
    h5 = doc.add_heading(level=1)
    h5_run = h5.add_run(f"5. Weighted Combined Course Outcome Attainment ({weighted.get('direct_weight', 80)}% Direct / {weighted.get('indirect_weight', 20)}% Indirect)")
    h5_run.font.color.rgb = RGBColor(30, 41, 59)

    wt_cos = weighted.get("combined_results", [])
    if wt_cos:
        table_wt = doc.add_table(rows=len(wt_cos) + 1, cols=7)
        hdrs = ["CO", "Direct %", "Direct Contrib", "Indirect %", "Indirect Contrib", "Combined %", "Combined Attainment"]
        for c_idx, h in enumerate(hdrs):
            cell = table_wt.rows[0].cells[c_idx]
            set_cell_background(cell, "1E293B")
            p = cell.paragraphs[0]
            run = p.add_run(h)
            run.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r_idx, c in enumerate(wt_cos, start=1):
            row = table_wt.rows[r_idx]
            vals = [
                c.get("co"), f"{c.get('direct_pct')}%", f"{c.get('direct_contribution')}%",
                f"{c.get('indirect_pct')}%", f"{c.get('indirect_contribution')}%",
                f"{c.get('combined_pct')}%", f"{c.get('combined_attainment')} ({c.get('level_label')})"
            ]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                if r_idx % 2 == 0:
                    set_cell_background(cell, "F8FAFC")
                p = cell.paragraphs[0]
                run = p.add_run(str(val))
                run.font.size = Pt(8.5)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != 6 else WD_ALIGN_PARAGRAPH.LEFT
        style_table(table_wt, [0.7, 0.9, 1.1, 0.9, 1.1, 1.0, 1.6])
    doc.add_paragraph()

    # Section 6: CO-PO Mapping & PO Attainment
    h6 = doc.add_heading(level=1)
    h6_run = h6.add_run("6. Program Outcome (PO & PSO) Attainment")
    h6_run.font.color.rgb = RGBColor(30, 41, 59)

    if po_results:
        has_indirect = copo.get("has_indirect") or any(p.get("indirect_attainment") is not None and p.get("indirect_attainment") > 0 for p in po_results)
        cols_count = 7 if has_indirect else 6
        table_po = doc.add_table(rows=len(po_results) + 1, cols=cols_count)
        if has_indirect:
            hdrs = ["Program Outcome", "Mapping Sum", "Avg Map", "Direct PO", "Indirect PO", "Final PO (0-3)", "Attainment Level"]
            col_widths = [1.1, 0.8, 0.8, 0.9, 0.9, 1.0, 1.3]
        else:
            hdrs = ["Program Outcome", "Mapping Sum", "Avg Map", "Direct PO (0-3)", "Final PO (0-3)", "Attainment Level"]
            col_widths = [1.3, 0.9, 0.9, 1.1, 1.1, 1.5]

        for c_idx, h in enumerate(hdrs):
            cell = table_po.rows[0].cells[c_idx]
            set_cell_background(cell, "1E293B")
            p = cell.paragraphs[0]
            run = p.add_run(h)
            run.bold = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        for r_idx, p_item in enumerate(po_results, start=1):
            row = table_po.rows[r_idx]
            att = p_item.get("attainment")
            att_disp = str(att) if att is not None else "N/A"
            dir_att = p_item.get("direct_attainment")
            dir_disp = str(dir_att) if dir_att is not None else "N/A"
            ind_att = p_item.get("indirect_attainment")
            ind_disp = str(ind_att) if ind_att is not None else "N/A"
            avg_m = str(p_item.get("mapping_avg", "N/A"))

            if has_indirect:
                vals = [
                    p_item.get("po"),
                    str(p_item.get("mapping_sum")),
                    avg_m,
                    dir_disp,
                    ind_disp,
                    att_disp,
                    p_item.get("level_label")
                ]
            else:
                vals = [
                    p_item.get("po"),
                    str(p_item.get("mapping_sum")),
                    avg_m,
                    dir_disp,
                    att_disp,
                    p_item.get("level_label")
                ]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                if r_idx % 2 == 0:
                    set_cell_background(cell, "F8FAFC")
                p = cell.paragraphs[0]
                run = p.add_run(str(val))
                run.font.size = Pt(8.5)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != len(vals) - 1 else WD_ALIGN_PARAGRAPH.LEFT
        style_table(table_po, col_widths)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out
